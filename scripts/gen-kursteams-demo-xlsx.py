"""Einmal-Skript: Demo-Kursteams in Import-Vorlage schreiben."""
import sys
from pathlib import Path

import openpyxl

src = Path(r'c:\Users\KurtSöser\Downloads\Kursteams-Import-Vorlage (1).xlsx')
out = Path(r'c:\Users\KurtSöser\Downloads\Kursteams-Import-Vorlage-DEMO.xlsx')

classes = ['1A', '1B', '2A', '2B', '3A', '4A', '4B']

plan = {
    '1A': [
        ('MAY', 'D'),
        ('GRO', 'E'),
        ('PAG', 'MAM'),
        ('TOW', 'MU'),
        ('COB', 'GES'),
        ('DEM1', 'MAM'),
        ('MAY', 'MU'),
    ],
    '1B': [
        ('COB', 'D'),
        ('MAY', 'E'),
        ('GRO', 'MAM'),
        ('PAG', 'MU'),
        ('TOW', 'GES'),
        ('DEM1', 'E'),
        ('COB', 'MAM'),
    ],
    '2A': [
        ('TOW', 'D'),
        ('COB', 'E'),
        ('MAY', 'MAM'),
        ('GRO', 'MU'),
        ('PAG', 'GES'),
        ('DEM1', 'GES'),
        ('TOW', 'MAM'),
    ],
    '2B': [
        ('DAL', 'D'),
        ('GRO', 'E'),
        ('TOW', 'MAM'),
        ('MAY', 'MU'),
        ('PAG', 'GES'),
        ('COB', 'E'),
        ('GRO', 'MAM'),
    ],
    '3A': [
        ('GRO', 'D'),
        ('PAG', 'E'),
        ('TOW', 'MAM'),
        ('COB', 'MU'),
        ('DEM1', 'GES'),
        ('MAY', 'MAM'),
        ('PAG', 'MU'),
    ],
    '4A': [
        ('PAG', 'D'),
        ('TOW', 'E'),
        ('COB', 'MAM'),
        ('MAY', 'MU'),
        ('GRO', 'GES'),
        ('DEM1', 'D'),
        ('TOW', 'GES'),
    ],
    '4B': [
        ('DEM1', 'D'),
        ('MAY', 'E'),
        ('PAG', 'MAM'),
        ('TOW', 'MU'),
        ('COB', 'GES'),
        ('GRO', 'E'),
        ('MAY', 'MU'),
    ],
}

rows = []
n = 1
for klasse in classes:
    for lehrer, fach in plan[klasse]:
        rows.append((f'demo-{n:03d}', lehrer, fach, klasse, ''))
        n += 1

rows.append(('demo-050', 'PAG', 'GES', '3A,4A', ''))

if not src.is_file():
    print(f'Quelle fehlt: {src}', file=sys.stderr)
    sys.exit(1)

wb = openpyxl.load_workbook(src)
ws = wb.active

if ws.max_row > 1:
    ws.delete_rows(2, ws.max_row - 1)

for i, row in enumerate(rows, start=2):
    for j, val in enumerate(row, start=1):
        ws.cell(row=i, column=j, value=val)

wb.save(out)
wb.save(src)
print(f'{len(rows)} Zeilen -> {out}')
print(f'auch aktualisiert: {src}')
